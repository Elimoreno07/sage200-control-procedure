from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from crear_control_procedure import (
    ErrorControlProcedure,
    detectar_ultima_version,
    generar_control_procedure,
    proponer_siguiente_version,
    validar_version,
)


class ControlProcedureTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporal = tempfile.TemporaryDirectory(dir=Path.cwd())
        self.carpeta = Path(self.temporal.name)

    def tearDown(self) -> None:
        self.temporal.cleanup()

    def _crear_libro(self, version: str) -> Path:
        ruta = self.carpeta / f"{version}.Control procedure.xlsx"
        libro = Workbook()
        principal = libro.active
        principal.title = "Control"
        principal["A1"] = f"Control {version} / {version}"
        principal["F1"] = "Resultado"
        principal["F2"] = "Eliminar"
        principal["G2"] = "Conservar"
        principal["H2"] = "=1+1"
        principal["F2"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        secundaria = libro.create_sheet("Datos")
        secundaria["B2"] = f"Versión {version}"
        secundaria["C2"] = f'="Versión {version}"'
        libro.save(ruta)
        libro.close()
        return ruta

    def test_validar_version(self) -> None:
        self.assertEqual(validar_version(" 2026.87.000 "), "2026.87.000")
        with self.assertRaises(ErrorControlProcedure):
            validar_version("2026-87-000")

    def test_proponer_siguiente_version(self) -> None:
        self.assertEqual(proponer_siguiente_version("2026.87.000"), "2026.87.001")
        self.assertEqual(proponer_siguiente_version("2026.87.999"), "2026.87.1000")

    def test_detectar_ultima_version_numerica(self) -> None:
        self._crear_libro("2026.9.999")
        ultima = self._crear_libro("2026.10.000")
        version, fichero = detectar_ultima_version(self.carpeta)
        self.assertEqual(version, "2026.10.000")
        self.assertEqual(fichero, ultima)

    def test_generacion_completa_preserva_formato_y_formulas(self) -> None:
        origen = self._crear_libro("2026.87.000")
        resultado = generar_control_procedure(
            "2026.87.000",
            "2026.90.000",
            self.carpeta,
            publicar_sharepoint=False,
        )

        self.assertTrue(resultado.destino.exists())
        self.assertEqual(resultado.origen, origen)
        self.assertEqual(resultado.hojas_modificadas, 2)
        self.assertEqual(resultado.reemplazos_realizados, 4)
        self.assertEqual(resultado.celdas_columna_f_limpiadas, 1)
        self.assertFalse(resultado.publicado_en_sharepoint)
        self.assertIsNone(resultado.destino_sharepoint)
        self.assertEqual(resultado.errores, ())
        self.assertTrue(resultado.log.exists())

        libro = load_workbook(resultado.destino, data_only=False)
        principal = libro["Control"]
        secundaria = libro["Datos"]
        self.assertEqual(principal["A1"].value, "Control 2026.90.000 / 2026.90.000")
        self.assertIsNone(principal["F2"].value)
        self.assertEqual(principal["F2"].fill.fgColor.rgb, "00FFFF00")
        self.assertEqual(principal["G2"].value, "Conservar")
        self.assertEqual(principal["H2"].value, "=1+1")
        self.assertEqual(secundaria["B2"].value, "Versión 2026.90.000")
        self.assertEqual(secundaria["C2"].value, '="Versión 2026.90.000"')
        libro.close()

        original = load_workbook(origen)
        self.assertEqual(original["Control"]["F2"].value, "Eliminar")
        self.assertIn("2026.87.000", original["Control"]["A1"].value)
        original.close()

    def test_no_sobrescribe_destino_existente(self) -> None:
        self._crear_libro("2026.87.000")
        existente = self._crear_libro("2026.90.000")
        contenido = existente.read_bytes()
        with self.assertRaises(ErrorControlProcedure):
            generar_control_procedure(
                "2026.87.000",
                "2026.90.000",
                self.carpeta,
                publicar_sharepoint=False,
            )
        self.assertEqual(existente.read_bytes(), contenido)

    def test_error_si_no_existe_origen(self) -> None:
        with self.assertRaisesRegex(ErrorControlProcedure, "No se ha encontrado"):
            generar_control_procedure(
                "2026.87.000",
                "2026.90.000",
                self.carpeta,
                publicar_sharepoint=False,
            )

    def test_publica_en_carpeta_sharepoint_sin_sobrescribir(self) -> None:
        self._crear_libro("2026.87.000")
        carpeta_sharepoint = self.carpeta / "sharepoint"
        carpeta_sharepoint.mkdir()

        resultado = generar_control_procedure(
            "2026.87.000",
            "2026.90.000",
            self.carpeta,
            publicar_sharepoint=True,
            carpeta_sharepoint=carpeta_sharepoint,
        )

        self.assertTrue(resultado.publicado_en_sharepoint)
        self.assertEqual(
            resultado.destino_sharepoint,
            carpeta_sharepoint / "2026.90.000.Control procedure.xlsx",
        )
        self.assertTrue(resultado.destino_sharepoint.exists())
        self.assertEqual(
            resultado.destino.read_bytes(), resultado.destino_sharepoint.read_bytes()
        )


if __name__ == "__main__":
    unittest.main()
