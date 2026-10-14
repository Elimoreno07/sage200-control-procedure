"""Lógica para generar versiones de Control Procedure de Sage 200."""

from __future__ import annotations

import logging
import os
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook


EXTENSIONES_EXCEL = {".xlsx", ".xlsm"}
PATRON_VERSION = re.compile(r"(?<!\d)(\d{4}\.\d{1,3}\.\d{3,4})(?!\d)")
PATRON_VERSION_COMPLETA = re.compile(r"^\d{4}\.\d{1,3}\.\d{3,4}$")
CARPETA_CONTROL_PROCEDURE = (
    Path.home() / "OneDrive - Sage Software, Inc" / "TechDoc - 2026"
)
CARPETA_SHAREPOINT = CARPETA_CONTROL_PROCEDURE
URL_SHAREPOINT = (
    "https://sage365.sharepoint.com/sites/global/epd/ib/TechDoc/"
    "Shared%20Documents/Forms/AllItems.aspx?"
    "id=%2Fsites%2Fglobal%2Fepd%2Fib%2FTechDoc%2FShared%20Documents%2F1%2ERelease%20Notes"
    "%2F4%2ESage%20200%2FCONTROL%2DPROCEDURE%2F2026&sortField=Modified&isAscending=false"
    "&viewid=8cf1bef6%2D5e1e%2D4247%2Db7aa%2Debecf7a34799&csf=1"
    "&CID=820d2693%2D48e6%2D41f4%2Db622%2D3aff44ff4ae0"
    "&FolderCTID=0x0120001201A6B66731EB4AB4006EB876706BF2"
)
ProgressCallback = Callable[[str, int], None]


class ErrorControlProcedure(Exception):
    """Error esperado y apto para mostrar a un usuario final."""


@dataclass(frozen=True)
class ResultadoGeneracion:
    """Resumen de una generación completada correctamente."""

    origen: Path
    destino: Path
    hojas_modificadas: int
    reemplazos_realizados: int
    celdas_columna_f_limpiadas: int
    publicado_en_sharepoint: bool
    destino_sharepoint: Path | None
    errores: tuple[str, ...]
    log: Path


def _notificar(callback: ProgressCallback | None, mensaje: str, porcentaje: int) -> None:
    if callback:
        callback(mensaje, porcentaje)


def _carpeta_datos_aplicacion() -> Path:
    base = Path(os.environ.get("LOCALAPPDATA", Path.home()))
    carpeta = base / "Sage200ControlProcedure"
    carpeta.mkdir(parents=True, exist_ok=True)
    return carpeta


def crear_logger() -> tuple[logging.Logger, Path]:
    """Crea un log independiente por ejecución."""
    carpeta_logs = _carpeta_datos_aplicacion() / "logs"
    carpeta_logs.mkdir(parents=True, exist_ok=True)
    fecha = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    ruta_log = carpeta_logs / f"control-procedure-{fecha}.log"

    logger = logging.getLogger(f"control_procedure.{fecha}")
    logger.setLevel(logging.INFO)
    logger.propagate = False
    handler = logging.FileHandler(ruta_log, encoding="utf-8")
    handler.setFormatter(
        logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    )
    logger.addHandler(handler)
    return logger, ruta_log


def validar_version(version: str) -> str:
    """Valida y normaliza una versión de Sage 200."""
    version = version.strip()
    if not PATRON_VERSION_COMPLETA.fullmatch(version):
        raise ErrorControlProcedure(
            f'La versión "{version}" no es válida. Usa un formato como 2026.87.000.'
        )
    return version


def extraer_version(nombre: str) -> str | None:
    coincidencia = PATRON_VERSION.search(nombre)
    return coincidencia.group(1) if coincidencia else None


def _clave_version(version: str) -> tuple[int, int, int]:
    partes = version.split(".")
    return int(partes[0]), int(partes[1]), int(partes[2])


def _misma_ruta(primera: Path, segunda: Path) -> bool:
    return os.path.normcase(str(primera.resolve())) == os.path.normcase(
        str(segunda.resolve())
    )


def listar_controles(carpeta: Path) -> list[tuple[tuple[int, int, int], str, Path]]:
    """Lista los Excel versionados de la carpeta, ordenados de menor a mayor."""
    if not carpeta.exists():
        raise ErrorControlProcedure(f"La carpeta seleccionada no existe:\n{carpeta}")
    if not carpeta.is_dir():
        raise ErrorControlProcedure(f"La ruta seleccionada no es una carpeta:\n{carpeta}")

    controles: list[tuple[tuple[int, int, int], str, Path]] = []
    try:
        for fichero in carpeta.iterdir():
            if (
                fichero.is_file()
                and fichero.suffix.lower() in EXTENSIONES_EXCEL
                and not fichero.name.startswith("~$")
            ):
                version = extraer_version(fichero.name)
                if version:
                    controles.append((_clave_version(version), version, fichero))
    except PermissionError as exc:
        raise ErrorControlProcedure(
            f"No tienes permiso para leer la carpeta seleccionada:\n{carpeta}"
        ) from exc
    except OSError as exc:
        raise ErrorControlProcedure(f"No se ha podido leer la carpeta: {exc}") from exc

    controles.sort(key=lambda elemento: (elemento[0], elemento[2].name.lower()))
    return controles


def detectar_ultima_version(carpeta: Path) -> tuple[str, Path]:
    """Devuelve la versión más alta y su fichero."""
    controles = listar_controles(carpeta)
    if not controles:
        raise ErrorControlProcedure(
            "No se ha encontrado ningún Control Procedure con una versión válida en:\n"
            f"{carpeta}"
        )

    clave_ultima = controles[-1][0]
    ultimos = [control for control in controles if control[0] == clave_ultima]
    if len(ultimos) > 1:
        nombres = "\n".join(f"  - {control[2].name}" for control in ultimos)
        raise ErrorControlProcedure(
            "Hay varios Excel para la última versión. Conserva un único fichero:\n"
            f"{nombres}"
        )
    _, version, fichero = ultimos[0]
    return version, fichero


def proponer_siguiente_version(version: str) -> str:
    """Incrementa el último bloque manteniendo su longitud mínima."""
    version = validar_version(version)
    anualidad, release, revision = version.split(".")
    siguiente = str(int(revision) + 1).zfill(len(revision))
    return f"{anualidad}.{release}.{siguiente}"


def buscar_fichero_origen(carpeta: Path, version_anterior: str) -> Path:
    """Localiza un único Excel cuya versión sea exactamente la solicitada."""
    version_anterior = validar_version(version_anterior)
    candidatos = [
        fichero
        for _, version, fichero in listar_controles(carpeta)
        if version == version_anterior
    ]

    if not candidatos:
        raise ErrorControlProcedure(
            f'No se ha encontrado ningún Excel para la versión "{version_anterior}" '
            f"en:\n{carpeta}"
        )
    if len(candidatos) > 1:
        listado = "\n".join(f"  - {fichero.name}" for fichero in candidatos)
        raise ErrorControlProcedure(
            "Se han encontrado varios Excel para la versión origen:\n" + listado
        )
    return candidatos[0]


def obtener_ultima_fila_con_datos(hoja) -> int:
    """Devuelve la última fila con valor, ignorando celdas solo formateadas."""
    for numero_fila in range(hoja.max_row, 0, -1):
        if any(celda.value is not None for celda in hoja[numero_fila]):
            return numero_fila
    return 1


def _actualizar_libro(
    fichero: Path,
    version_anterior: str,
    nueva_version: str,
    logger: logging.Logger,
) -> tuple[int, int, int]:
    """Actualiza versiones y limpia F en la hoja activa sin alterar estilos."""
    conservar_macros = fichero.suffix.lower() == ".xlsm"
    try:
        libro = load_workbook(fichero, keep_vba=conservar_macros, keep_links=True)
    except PermissionError as exc:
        raise ErrorControlProcedure(
            f'No se puede abrir "{fichero.name}". Comprueba que Excel esté cerrado.'
        ) from exc
    except Exception as exc:
        raise ErrorControlProcedure(
            f'No se ha podido abrir "{fichero.name}". Puede estar dañado: {exc}'
        ) from exc

    hojas_modificadas: set[str] = set()
    reemplazos = 0
    celdas_f_limpiadas = 0

    try:
        for hoja in libro.worksheets:
            reemplazos_hoja = 0
            for fila in hoja.iter_rows():
                for celda in fila:
                    valor = celda.value
                    if isinstance(valor, str) and version_anterior in valor:
                        apariciones = valor.count(version_anterior)
                        celda.value = valor.replace(version_anterior, nueva_version)
                        reemplazos += apariciones
                        reemplazos_hoja += apariciones
            if reemplazos_hoja:
                hojas_modificadas.add(hoja.title)
                logger.info(
                    "Hoja %s: %s reemplazos de versión", hoja.title, reemplazos_hoja
                )

        hoja_activa = libro.active
        ultima_fila = obtener_ultima_fila_con_datos(hoja_activa)
        for numero_fila in range(2, ultima_fila + 1):
            celda = hoja_activa.cell(row=numero_fila, column=6)
            if celda.value is not None:
                celda.value = None
                celdas_f_limpiadas += 1
        if celdas_f_limpiadas:
            hojas_modificadas.add(hoja_activa.title)
        logger.info(
            "Columna F limpiada en %s: %s celdas",
            hoja_activa.title,
            celdas_f_limpiadas,
        )

        try:
            libro.save(fichero)
        except PermissionError as exc:
            raise ErrorControlProcedure(
                f'No se puede guardar "{fichero.name}". Cierra el fichero en Excel '
                "y vuelve a intentarlo."
            ) from exc
        except OSError as exc:
            raise ErrorControlProcedure(f"No se ha podido guardar el Excel: {exc}") from exc
    finally:
        libro.close()

    return len(hojas_modificadas), reemplazos, celdas_f_limpiadas


def generar_control_procedure(
    version_anterior: str,
    nueva_version: str,
    carpeta: Path,
    progress_callback: ProgressCallback | None = None,
    publicar_sharepoint: bool = True,
    carpeta_sharepoint: Path | None = None,
) -> ResultadoGeneracion:
    """Genera un Control Procedure de forma transaccional y devuelve su resumen."""
    logger, ruta_log = crear_logger()
    temporal: Path | None = None
    destino: Path | None = None
    destino_sharepoint: Path | None = None
    destino_creado_por_proceso = False
    sharepoint_creado_por_proceso = False
    try:
        version_anterior = validar_version(version_anterior)
        nueva_version = validar_version(nueva_version)
        if version_anterior == nueva_version:
            raise ErrorControlProcedure(
                "La versión origen y la versión destino deben ser diferentes."
            )

        logger.info("Inicio de generación")
        logger.info("Carpeta de trabajo: %s", carpeta)
        logger.info("Versión origen: %s", version_anterior)
        logger.info("Versión destino: %s", nueva_version)

        _notificar(progress_callback, "Buscando el fichero origen...", 10)
        origen = buscar_fichero_origen(carpeta, version_anterior)
        logger.info("Fichero origen: %s", origen)

        nombre_destino = origen.name.replace(version_anterior, nueva_version)
        if nombre_destino == origen.name:
            raise ErrorControlProcedure(
                "No se puede generar el nombre destino porque el nombre del origen no "
                "contiene su versión."
            )
        destino = origen.with_name(nombre_destino)
        if destino.exists():
            raise ErrorControlProcedure(
                f"El fichero destino ya existe y no se sobrescribirá:\n{destino}"
            )

        if publicar_sharepoint:
            carpeta_publicacion = carpeta_sharepoint or CARPETA_SHAREPOINT
            if not carpeta_publicacion.is_dir():
                raise ErrorControlProcedure(
                    "No se encuentra el acceso sincronizado de SharePoint:\n"
                    f"{carpeta_publicacion}\n\n"
                    "Añade el acceso directo 'TechDoc - 2026' a OneDrive o "
                    "selecciónalo como carpeta de trabajo."
                )
            destino_sharepoint = carpeta_publicacion / nombre_destino
            if not _misma_ruta(destino, destino_sharepoint) and destino_sharepoint.exists():
                raise ErrorControlProcedure(
                    "El fichero ya existe en SharePoint y no se sobrescribirá:\n"
                    f"{destino_sharepoint}"
                )

        _notificar(progress_callback, "Creando una copia de trabajo...", 30)
        descriptor, nombre_temporal = tempfile.mkstemp(
            prefix="control-procedure-", suffix=origen.suffix.lower()
        )
        os.close(descriptor)
        temporal = Path(nombre_temporal)
        shutil.copy2(origen, temporal)

        _notificar(progress_callback, "Actualizando el contenido del Excel...", 55)
        hojas, reemplazos, celdas_f = _actualizar_libro(
            temporal, version_anterior, nueva_version, logger
        )

        _notificar(progress_callback, "Guardando el nuevo fichero...", 85)
        try:
            destino_creado_por_proceso = True
            shutil.copy2(temporal, destino)
        except PermissionError as exc:
            raise ErrorControlProcedure(
                "No se ha podido guardar el fichero destino. Comprueba que Excel esté "
                "cerrado y que tengas permisos en la carpeta."
            ) from exc
        except OSError as exc:
            raise ErrorControlProcedure(f"No se ha podido crear el destino: {exc}") from exc

        logger.info("Fichero destino: %s", destino)
        if publicar_sharepoint and destino_sharepoint:
            _notificar(progress_callback, "Publicando en SharePoint...", 95)
            if not _misma_ruta(destino, destino_sharepoint):
                sharepoint_creado_por_proceso = True
                try:
                    shutil.copy2(destino, destino_sharepoint)
                except PermissionError as exc:
                    raise ErrorControlProcedure(
                        "No se ha podido publicar en SharePoint. Comprueba que OneDrive "
                        "esté iniciado y que tengas permisos en 'TechDoc - 2026'."
                    ) from exc
                except OSError as exc:
                    raise ErrorControlProcedure(
                        f"No se ha podido publicar en SharePoint: {exc}"
                    ) from exc
            logger.info("Publicado mediante OneDrive en: %s", destino_sharepoint)
            logger.info("Carpeta web: %s", URL_SHAREPOINT)
        logger.info("Hojas modificadas: %s", hojas)
        logger.info("Reemplazos realizados: %s", reemplazos)
        logger.info("Generación completada correctamente")
        _notificar(progress_callback, "Proceso completado.", 100)
        return ResultadoGeneracion(
            origen=origen,
            destino=destino,
            hojas_modificadas=hojas,
            reemplazos_realizados=reemplazos,
            celdas_columna_f_limpiadas=celdas_f,
            publicado_en_sharepoint=publicar_sharepoint,
            destino_sharepoint=destino_sharepoint,
            errores=(),
            log=ruta_log,
        )
    except ErrorControlProcedure as exc:
        logger.error("%s", exc)
        if sharepoint_creado_por_proceso and destino_sharepoint and destino_sharepoint.exists():
            try:
                destino_sharepoint.unlink()
            except OSError:
                logger.exception("No se pudo retirar la publicación incompleta")
        if destino_creado_por_proceso and destino and destino.exists():
            try:
                destino.unlink()
            except OSError:
                logger.exception("No se pudo retirar el destino incompleto")
        raise ErrorControlProcedure(f"{exc}\n\nLog: {ruta_log}") from exc
    except Exception as exc:
        logger.exception("Error inesperado")
        if sharepoint_creado_por_proceso and destino_sharepoint and destino_sharepoint.exists():
            try:
                destino_sharepoint.unlink()
            except OSError:
                logger.exception("No se pudo retirar la publicación incompleta")
        if destino_creado_por_proceso and destino and destino.exists():
            try:
                destino.unlink()
            except OSError:
                logger.exception("No se pudo retirar el destino incompleto")
        raise ErrorControlProcedure(
            f"Se ha producido un error inesperado: {exc}\nConsulta el log: {ruta_log}"
        ) from exc
    finally:
        if temporal and temporal.exists():
            try:
                temporal.unlink()
            except OSError:
                logger.exception("No se pudo eliminar el fichero temporal")
        for handler in list(logger.handlers):
            handler.close()
            logger.removeHandler(handler)


def crear_control_procedure(
    version_anterior: str, nueva_version: str, carpeta: Path
) -> Path:
    """API compatible con versiones anteriores del proyecto."""
    return generar_control_procedure(version_anterior, nueva_version, carpeta).destino


def pedir_dato(mensaje: str) -> str:
    valor = input(mensaje).strip()
    if not valor:
        raise ErrorControlProcedure("El valor indicado no puede estar vacío.")
    return valor


def main() -> int:
    """Entrada alternativa por consola."""
    print("Creación de Control Procedure para Sage 200")
    print("-" * 44)
    try:
        version_anterior = pedir_dato("Versión origen (ej. 2026.87.000): ")
        nueva_version = pedir_dato("Versión destino (ej. 2026.90.000): ")
        resultado = generar_control_procedure(
            version_anterior, nueva_version, CARPETA_CONTROL_PROCEDURE
        )
    except (ErrorControlProcedure, EOFError, KeyboardInterrupt) as exc:
        mensaje = "Operación cancelada." if isinstance(exc, KeyboardInterrupt) else str(exc)
        print(f"\nERROR: {mensaje}", file=sys.stderr)
        return 1

    print(f"Archivo origen: {resultado.origen}")
    print(f"Archivo destino: {resultado.destino}")
    print(f"Hojas modificadas: {resultado.hojas_modificadas}")
    print(f"Reemplazos: {resultado.reemplazos_realizados}")
    print(f"Publicado en SharePoint: {'Sí' if resultado.publicado_en_sharepoint else 'No'}")
    if resultado.destino_sharepoint:
        print(f"Destino SharePoint: {resultado.destino_sharepoint}")
    print(f"Log: {resultado.log}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
